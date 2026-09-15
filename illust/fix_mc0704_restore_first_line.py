#!/usr/bin/env python3
"""v60.6 — MC-7-04-001 의 ai_first_line 을 기획자가 원래 정한 값으로 되돌린다.

v60.5 에서 스토리를 강아지 돌봄으로 바꾸며 ai_first_line 자체까지 새로
썼는데, 이건 넘겨짚은 것이었다 — 기획자는 스토리(누구에게 왜 부탁하는가)만
바꿔 달라고 했지 첫 대사를 바꿔 달라고 한 적이 없다("난 첫 대사 바꾸라고 한
적은 없어", 2026-09-15). 첫 대사는 기획자가 그전에 이미 직접
「내일 바쁘세요? 급히 도움이 필요해서요.」로 정해 둔 값이었다(v60.2 는 이
값을 존중해 그대로 두고, 사정 공개만 한 턴 늦추는 지연 공개 장치를 붙였다).

되돌리는 것
    - ai_first_line: 새로 쓴 문장 → 기획자가 정한 원래 문장
    - "# 첫 대사 (고정)" 블록: v60.2 의 지연 공개 구조를 복원하되,
      공개되는 내용만 이사→강아지로 바꾼다.
    - "대화의 흐름": 마찬가지로 지연 공개 구조를 복원하고 내용만 강아지로.
미션 3개의 예시 문장(v60.5 에서 이미 강아지 스토리로 바꿔 둔 것)은 그대로
둔다 — 어차피 지연 공개 뒤(사정을 이미 들은 뒤)에 나오는 반응이라 구조가
바뀌어도 그대로 맞는다.

돌리기
    python3 illust/fix_mc0704_restore_first_line.py            # 확인만
    python3 illust/fix_mc0704_restore_first_line.py --write    # 실제 반영
"""
from __future__ import annotations

import argparse
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
LEDGER = ROOT / "글로벌_교재기반_콘텐츠_v60.xlsx"

WRONG_FIRST_LINE = (
    "미안한데 제가 갑자기 급한 일이 생겨서 그런데, 혹시 내일 하루만 저희 "
    "강아지 좀 봐주실 수 있어요?"
)
CORRECT_FIRST_LINE = "내일 바쁘세요? 급히 도움이 필요해서요."
REVEAL_LINE = (
    "사실 제가 갑자기 급한 일이 생겨서 그런데, 혹시 내일 하루만 저희 강아지 "
    "좀 봐주실 수 있어요?"
)

WRONG_FIRST_LINE_BLOCK = (
    "# 첫 대사 (고정)\n"
    "대화는 반드시 아래 문장으로 시작합니다. 이 문장만 말하고 사용자의 응답을 기다리세요.\n"
    f"「{WRONG_FIRST_LINE}」\n"
    "※ 이 문장은 고정 대사입니다. 아래의 어떤 규칙보다 우선하며 한 글자도 바꾸지 마세요.\n"
    "※ 두 번째 문장부터는 사용자의 응답을 듣고 나서 말합니다."
)
CORRECT_FIRST_LINE_BLOCK = (
    "# 첫 대사 (고정)\n"
    "대화는 반드시 아래 문장으로 시작합니다. 이 문장만 말하고 사용자의 응답을 기다리세요.\n"
    f"「{CORRECT_FIRST_LINE}」\n"
    "※ 이 문장은 고정 대사입니다. 아래의 어떤 규칙보다 우선하며 한 글자도 바꾸지 마세요.\n"
    "※ 두 번째 문장부터는 사용자의 응답을 듣고 나서 말합니다.\n"
    f"※ 사용자가 무슨 일인지 물으면 그때 사정을 설명하세요: 「{REVEAL_LINE}」"
)

WRONG_FLOW = (
    "AI는 갑자기 급한 일이 생겼다며 내일 하루만 반려견을 봐 줄 수 있는지 "
    "묻는다. 사용자는 그 부탁에 놀라고, 상대의 사정에 공감한 뒤, 자신도 "
    "도저히 봐 줄 수 없는 이유를 들어 정중하게 거절한다. 대화는 처음부터 "
    "끝까지 존댓말로 이어진다."
)
CORRECT_FLOW = (
    "AI는 내일 바쁜지 물으며 급히 도움이 필요하다고만 말한다 — 이 시점에는 "
    "무슨 일인지 밝히지 않는다. 사용자가 무슨 일이냐고 물으면 그제서야 갑자기 "
    "급한 일이 생겨서 내일 하루만 반려견을 봐 줄 수 있는지 묻는다. 사용자는 "
    "그 부탁에 놀라고, 상대의 사정에 공감한 뒤, 자신도 도저히 봐 줄 수 없는 "
    "이유를 들어 정중하게 거절한다. 대화는 처음부터 끝까지 존댓말로 이어진다."
)

CHANGE_NOTE = (
    "v60.6(2026-09-15) v60.5 되돌림 — 기획자는 스토리만 바꿔 달라고 했지 "
    "첫 대사를 바꿔 달라고 한 적이 없었다. ai_first_line 을 기획자가 원래 "
    "정한 「내일 바쁘세요? 급히 도움이 필요해서요.」로 복원하고, 강아지 사정은 "
    "v60.2 와 같은 지연 공개 구조(사용자가 물으면 그때 밝힌다)로 다시 붙임"
)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--write", action="store_true")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(LEDGER)
    ws = wb["n7_mission_chat"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}

    touched = 0
    for row in ws.iter_rows(min_row=2):
        if row[col["item_id"] - 1].value != "MC-7-04-001":
            continue
        prompt_cell = row[col["ai_persona_prompt"] - 1]
        first_line_cell = row[col["ai_first_line"] - 1]
        prompt = prompt_cell.value

        assert first_line_cell.value == WRONG_FIRST_LINE, "ai_first_line 이 예상과 다르다"
        assert WRONG_FIRST_LINE_BLOCK in prompt, "첫 대사 블록을 못 찾았다"
        assert WRONG_FLOW in prompt, "대화의 흐름 문장을 못 찾았다"

        new_prompt = prompt.replace(WRONG_FIRST_LINE_BLOCK, CORRECT_FIRST_LINE_BLOCK)
        new_prompt = new_prompt.replace(WRONG_FLOW, CORRECT_FLOW)

        print("=== 복원된 첫 대사 ===")
        print(CORRECT_FIRST_LINE)
        print("=== 지연 공개 문장 ===")
        print(REVEAL_LINE)
        print("=== 복원된 대화 흐름 ===")
        print(CORRECT_FLOW)

        if args.write:
            first_line_cell.value = CORRECT_FIRST_LINE
            prompt_cell.value = new_prompt
            cn_cell = row[col["change_note"] - 1]
            old_note = cn_cell.value
            cn_cell.value = f"{old_note} / {CHANGE_NOTE}" if old_note else CHANGE_NOTE
        touched += 1
    assert touched == 1, f"MC-7-04-001 을 정확히 한 번 고쳐야 하는데 {touched}번 걸렸다"

    if not args.write:
        print("\n--dry(기본값)라 원장을 쓰지 않았다. --write 로 반영해라.")
        return

    backup = LEDGER.with_name(
        LEDGER.stem + f".backup-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    )
    shutil.copy2(LEDGER, backup)
    wb.save(LEDGER)
    print(f"\n저장했다. 백업 — {backup.name}")


if __name__ == "__main__":
    main()
